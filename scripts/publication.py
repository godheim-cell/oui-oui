#!/usr/bin/env python3
"""Construction générique des livrables Nexus à partir d'un volume manifesté."""
from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import tomllib
import unicodedata
from pathlib import Path


def load_config(root: Path) -> dict:
    with (root / "config/projet.toml").open("rb") as handle:
        return tomllib.load(handle)


def safe_filename(value: str) -> str:
    ascii_value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^A-Za-z0-9._-]+", "_", ascii_value).strip("_") or "manuscrit"


def volume_paths(root: Path, volume: str) -> list[Path]:
    from lifecycle import volume_paths as lifecycle_volume_paths
    return lifecycle_volume_paths(root, volume)


def visible_scene(root: Path, path: Path) -> tuple[dict, str]:
    import nexus
    old_root = nexus.ROOT
    nexus.ROOT = root
    try:
        return nexus.frontmatter(path)
    finally:
        nexus.ROOT = old_root


def output_basename(root: Path, volume: str) -> str:
    cfg = load_config(root)
    return safe_filename(f"{cfg['project'].get('title', 'Manuscrit')}_{volume}")


def build_markdown(root: Path, volume: str) -> Path:
    cfg = load_config(root)
    paths = volume_paths(root, volume)
    directory = root / "publication" / volume
    directory.mkdir(parents=True, exist_ok=True)
    chunks = [f"# {cfg['project']['title']}", "", f"**Auteur :** {cfg['project'].get('author', 'unknown')}", "", f"**Volume :** {volume}", ""]
    last_chapter = None
    for path in paths:
        meta, body = visible_scene(root, path)
        chapter = meta.get("chapter")
        if chapter != last_chapter:
            chunks += [f"## Chapitre {chapter}", ""]
            last_chapter = chapter
        chunks += [body.strip(), "", "---", ""]
    target = directory / f"{output_basename(root, volume)}.md"
    target.write_text("\n".join(chunks).rstrip() + "\n", encoding="utf-8")
    return target


def _add_body_paragraphs(document, body: str) -> None:
    from docx.shared import Cm
    for block in re.split(r"\n\s*\n", body.strip()):
        text = block.strip()
        if not text:
            continue
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = 0
        if not text.startswith("—"):
            paragraph.paragraph_format.first_line_indent = Cm(0.5)
        paragraph.add_run(text)


def build_docx(root: Path, volume: str) -> Path:
    try:
        from docx import Document
        from docx.shared import Cm, Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError as exc:
        raise RuntimeError("python-docx est requis pour générer le DOCX.") from exc

    cfg = load_config(root)
    pub = cfg.get("publication", {})
    paths = volume_paths(root, volume)
    directory = root / "publication" / volume
    directory.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(float(pub.get("page_width_cm", 14.8)))
    section.page_height = Cm(float(pub.get("page_height_cm", 21.0)))
    section.top_margin = Cm(float(pub.get("margin_top_cm", 1.7)))
    section.bottom_margin = Cm(float(pub.get("margin_bottom_cm", 1.7)))
    section.left_margin = Cm(float(pub.get("margin_left_cm", 1.8)))
    section.right_margin = Cm(float(pub.get("margin_right_cm", 1.8)))

    document.styles["Normal"].font.name = str(pub.get("font", "Liberation Serif"))
    document.styles["Normal"].font.size = Pt(float(pub.get("font_size_pt", 10.5)))

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(cfg["project"]["title"])
    run.bold = True
    run.font.size = Pt(20)
    author = document.add_paragraph()
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author.add_run(cfg["project"].get("author", "unknown"))
    volume_p = document.add_paragraph()
    volume_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    volume_p.add_run(volume)
    document.add_page_break()

    last_chapter = None
    for path in paths:
        meta, body = visible_scene(root, path)
        chapter = meta.get("chapter")
        if chapter != last_chapter:
            if last_chapter is not None:
                document.add_page_break()
            heading = document.add_paragraph()
            heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
            hrun = heading.add_run(f"Chapitre {chapter}")
            hrun.bold = True
            hrun.font.size = Pt(14)
            last_chapter = chapter
        _add_body_paragraphs(document, body)

    props = document.core_properties
    props.title = cfg["project"]["title"]
    props.author = cfg["project"].get("author", "unknown")
    props.subject = f"{cfg['project'].get('series', '')} — {volume}".strip(" —")
    props.comments = "Généré depuis les scènes manifestées par Nexus."

    target = directory / f"{output_basename(root, volume)}.docx"
    document.save(target)
    return target


def build_pdf(root: Path, docx_path: Path) -> Path:
    executable = shutil.which("soffice") or shutil.which("libreoffice")
    if not executable:
        raise RuntimeError("LibreOffice/soffice est requis pour convertir le DOCX en PDF.")
    outdir = docx_path.parent
    subprocess.run([executable, "--headless", "--convert-to", "pdf", "--outdir", str(outdir), str(docx_path)], check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
    target = outdir / f"{docx_path.stem}.pdf"
    if not target.is_file() or target.stat().st_size == 0:
        raise RuntimeError("La conversion PDF n'a pas produit de fichier valide.")
    return target


def build_publication(root: Path, volume: str, want_docx: bool = True, want_pdf: bool = False) -> list[Path]:
    from lifecycle import publication_manifest
    outputs = [build_markdown(root, volume)]
    docx = None
    if want_docx or want_pdf:
        docx = build_docx(root, volume)
        outputs.append(docx)
    if want_pdf:
        outputs.append(build_pdf(root, docx))
    outputs.append(publication_manifest(root, volume))
    return outputs


def main() -> int:
    root = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(prog="nexus-publication")
    parser.add_argument("--volume", required=True)
    parser.add_argument("--no-docx", action="store_true")
    parser.add_argument("--pdf", action="store_true")
    args = parser.parse_args()
    for path in build_publication(root, args.volume, want_docx=not args.no_docx, want_pdf=args.pdf):
        print(path.relative_to(root))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
